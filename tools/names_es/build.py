#!/usr/bin/env python3
"""Build es/names/{given_names,surnames,places}.tsv from official open data.

Sources (download with tools/names_es/fetch.sh into RAW_DIR; licences in SOURCES.md):
  INE, nombres por frecuencia (Censos anuales 2025)   (CC BY 4.0)  given names
  INE, apellidos por frecuencia (Censos anuales 2025)  (CC BY 4.0)  surnames
  INE, Padron municipal 2025 (cifras por municipio)    (CC BY 4.0)  municipalities, provinces
  GeoNames extract (geonames_extract.py)               (CC BY 4.0)  foreign countries, cities, regions
plus our own curation files in es/names/curated/ (see es/names/README.md).

INE writes person names in capitals without accents (JOSE, GONZALEZ); the usual spelling comes
from curated/spellings.tsv. Usage: build.py RAW_DIR GEO_TSV OUT_DIR [--stats FILE]
Needs openpyxl and xlrd.
"""
import re
import sys
import unicodedata
from collections import defaultdict

import openpyxl
import xlrd

# ---- thresholds (documented in es/names/README.md) ----
GIVEN_MIN_PERSONS = 100       # residents bearing the name (alone or as part of a compound name)
GENDER_SHARE = 0.90           # m (f) if at least this share of the bearers are men (women), else u
SURNAME_MIN_PERSONS = 200     # residents whose first surname is or contains this name
MUNICIPALITY_MIN_POP = 500    # Padron 2025
FOREIGN_CITY_MIN_POP = 150_000
FOREIGN_ADM1_MIN_POP = 1_000_000

GIVEN_FILE = "ine/nombres_por_edad_media.xlsx"
SURNAME_FILE = "ine/apellidos_frecuencia.xls"
POP_FILE = "mun/pobmun25.xlsx"

# Articles (Castilian, Catalan, Galician, Balearic) and their gender, postposed by INE
# (`Coruña, A`, `Palmas, Las`, `Alfàs del Pi, l'`).
ARTICLES = {
    "el": "m", "los": "m", "lo": "m", "els": "m", "o": "m", "os": "m", "es": "m",
    "la": "f", "las": "f", "les": "f", "a": "f", "as": "f", "sa": "f", "ses": "f",
    "l'": "u", "s'": "u",
}
# Particles and articles: never listed as name parts.
PARTICLES = {
    "el", "la", "los", "las", "lo", "l", "els", "les", "o", "os", "a", "as", "es", "sa", "ses",
    "s", "de", "del", "d", "dels", "y", "i", "e", "en", "da", "do", "das", "dos", "di", "al",
    "san", "santa", "santo", "sant", "sta", "sto", "van", "von", "der", "den", "le", "du",
    "des", "ben", "bin", "abu", "ibn", "mc", "mac", "and", "of", "the", "upon", "on",
}


def fold(s):
    return "".join(c for c in unicodedata.normalize("NFD", s) if not unicodedata.combining(c)).lower()


def sort_key(name):
    return (fold(name), name)


def title(upper):
    """INE capitals -> usual case per word part: JOSE -> Jose, MARIA-JOSE -> Maria-Jose."""
    out, start = [], True
    for c in upper:
        out.append(c.upper() if start else c.lower())
        start = c in " -'"
    return "".join(out)


def ok_name(n):
    """At least 3 letters, nothing but letters and inner hyphens (no apostrophe, middle dot,
    period or digit: the tokenizer splits there)."""
    return len(n) >= 3 and n.replace("-", "").isalpha() and not n.startswith("-") and not n.endswith("-")


WORDS = set()      # names that are also ordinary Spanish words (curated/word.tsv)
EXCLUDED = defaultdict(set)  # name -> lists it is left out of (given, surname, place)


def parts(name, lst="place"):
    """Whole name (if one token, possibly hyphenated) and its name parts.

    A part that is an ordinary word is left out (`Frontera` of Jerez de la Frontera); a whole
    name that is one is kept and marked."""
    name = name.replace("’", "'").strip()
    if "," in name or "(" in name:  # `Noáin (Valle de Elorz)/Noain (Elortzibar)`
        out = []
        for piece in name.replace("(", ",").replace(")", ",").split(","):
            if piece.strip():
                out += parts(piece.strip(), lst)
        return list(dict.fromkeys(out))
    out = []
    tokens = name.split()

    def whole(tok):
        segs = tok.split("-")
        return ok_name(tok) and segs[0].lower() not in PARTICLES and segs[-1].lower() not in PARTICLES

    if len(tokens) == 1 and whole(name):
        out.append(name)
    for tok in tokens:
        if "-" in tok and len(tokens) > 1 and whole(tok):
            out.append(tok)
        for p in tok.split("-"):
            if "'" in p:  # Catalan elision: d'Alcalà, l'Hospitalet -> Alcalà, Hospitalet
                head, _, p = p.partition("'")
                if head.lower() not in ("d", "l", "s"):
                    continue
            if p.lower() in PARTICLES or not ok_name(p) or p[0].islower():
                continue
            if p != name and (p in WORDS or lst in EXCLUDED[p]):
                continue
            out.append(p)
    return [n for n in dict.fromkeys(out) if lst not in EXCLUDED[n]]


def read_curated(path, ncols):
    rows = []
    try:
        with open(path, encoding="utf-8") as f:
            for line in f:
                if line.strip():
                    p = line.rstrip("\n").split("\t")
                    rows.append(p + [""] * (ncols - len(p)))
    except FileNotFoundError:
        pass
    return rows


def ine_place(name):
    """INE municipality or province name -> [(name without article, gender)] per language form.

    `Alacant/Alicante` -> two names; `Coruña, A` -> (`Coruña`, f); `Palmas, Las` -> (`Palmas`, f)."""
    out = []
    for alt in name.split("/"):
        alt = alt.strip()
        g = "u"
        m = re.match(r"^(.*), *(\S+)$", alt)
        if m and m.group(2).lower() in ARTICLES:
            alt, g = m.group(1), ARTICLES[m.group(2).lower()]
        out.append((alt, g))
    return out


def main(raw, geo, out, stats_path=None):
    cur = f"{out}/curated"
    words = {r[0] for r in read_curated(f"{cur}/word.tsv", 1)}
    WORDS.update(words)
    for r in read_curated(f"{cur}/excluded.tsv", 3):
        EXCLUDED[r[0]].add(r[1])
    spell = {}  # (list, register) -> [spellings]
    for r in read_curated(f"{cur}/spellings.tsv", 4):
        spell[(r[1], r[0])] = [s for s in r[2:] if s]
    stats = []

    def spellings(lst, reg):
        return spell.get((lst, reg), [title(reg)])

    # ---- given names ----
    # counts[name] = [men alone or first, women alone or first, men any position, women any]
    counts = defaultdict(lambda: [0, 0, 0, 0])
    persons, rows_given = [0, 0], []
    wb = openpyxl.load_workbook(f"{raw}/{GIVEN_FILE}", read_only=True)
    for sex, ws in enumerate(wb.worksheets[:2]):  # Hombres, Mujeres
        for r in list(ws.iter_rows(values_only=True))[7:]:
            if not r[1] or not isinstance(r[2], (int, float)):
                continue
            toks = r[1].split()
            persons[sex] += r[2]
            rows_given.append((toks, r[2]))
            for i, t in enumerate(toks):
                if i == 0:
                    counts[t][sex] += r[2]
                counts[t][2 + sex] += r[2]
    given = {}
    listed_reg = set()
    for reg, (m1, f1, ma, fa) in counts.items():
        if ma + fa < GIVEN_MIN_PERSONS or reg.lower() in PARTICLES:
            continue
        m, f = (m1, f1) if m1 + f1 else (ma, fa)
        g = "m" if m >= GENDER_SHARE * (m + f) else "f" if f >= GENDER_SHARE * (m + f) else "u"
        got = False
        for n in spellings("given", reg):
            if not ok_name(n) or "given" in EXCLUDED[n] or "given" in EXCLUDED[title(reg)]:
                continue
            given[n] = g if given.get(n, g) == g else "u"
            got = True
        if got:
            listed_reg.add(reg)
    cov = sum(c for toks, c in rows_given if all(t in listed_reg or t.lower() in PARTICLES for t in toks))
    stats.append(f"given names: {len(given)} (>= {GIVEN_MIN_PERSONS} bearers), "
                 f"all given names listed for {cov / sum(persons):.1%} of the residents in the "
                 f"INE file; m {sum(1 for g in given.values() if g == 'm')}, "
                 f"f {sum(1 for g in given.values() if g == 'f')}, "
                 f"u {sum(1 for g in given.values() if g == 'u')}")

    # ---- surnames ----
    first = defaultdict(float)   # surname part -> residents with it in the first surname
    rows_sur = []
    wb = xlrd.open_workbook(f"{raw}/{SURNAME_FILE}")
    for ws in wb.sheets():
        for i in range(5, ws.nrows):
            v = ws.row_values(i)
            if not v[1] or not isinstance(v[2], float):
                continue
            toks = [t for t in v[1].split() if t.lower() not in PARTICLES]
            rows_sur.append((toks, v[2]))
            for t in toks:
                first[t] += v[2]
    surnames, listed_reg = set(), set()
    for reg, c in first.items():
        if c < SURNAME_MIN_PERSONS:
            continue
        for n in spellings("surname", reg):
            if ok_name(n) and "surname" not in EXCLUDED[n] and "surname" not in EXCLUDED[title(reg)]:
                surnames.add(n)
                listed_reg.add(reg)
    total = sum(c for _, c in rows_sur)
    cov = sum(c for toks, c in rows_sur if toks and all(t in listed_reg for t in toks))
    stats.append(f"surnames: {len(surnames)} (>= {SURNAME_MIN_PERSONS} bearers of the first "
                 f"surname), first surname listed for {cov / total:.1%} of the residents in the "
                 f"INE file (first surnames with >= 20 bearers)")

    # ---- places ----
    places = {}

    def add_place(name, gender, kind):
        for n in parts(name):
            is_whole = n == name
            g = gender if is_whole else "u"
            if n not in places:
                places[n] = (g, kind, is_whole)
                continue
            og, ok, ow = places[n]
            if is_whole and not ow:  # the name itself beats a part of a longer name
                places[n] = (g, kind, True)
            elif is_whole == ow:
                ng = g if og == "u" else og if g in ("u", og) else "u"
                places[n] = (ng, ok, ow)

    for row in read_curated(f"{cur}/places.tsv", 3):
        add_place(row[0], row[1] or "u", row[2])
    rows = [r for r in list(openpyxl.load_workbook(f"{raw}/{POP_FILE}", read_only=True)
                            .worksheets[0].iter_rows(values_only=True))[2:] if r[0] and r[4]]
    for prov in sorted({r[1] for r in rows}):
        for n, g in ine_place(prov):
            add_place(n, g, "state")
    pop_all = sum(r[4] for r in rows)
    n_mun = pop_cov = 0
    for r in rows:
        if r[4] < MUNICIPALITY_MIN_POP:
            continue
        n_mun += 1
        pop_cov += r[4]
        for n, g in ine_place(r[3]):
            add_place(n, g, "city")
    stats.append(f"municipalities: {n_mun} with >= {MUNICIPALITY_MIN_POP} inhabitants, "
                 f"{pop_cov / pop_all:.1%} of the population of Spain")
    n_city = n_adm = n_country = 0
    with open(geo, encoding="utf-8") as f:
        for line in f:
            _gid, fc, cc, p, _n, esname, _name = line.rstrip("\n").split("\t")
            p = int(p)
            if fc.startswith("A.PCL") or fc == "L.CONT":
                n_country += 1
                add_place(esname, "u", "country" if fc.startswith("A.PCL") else "continent")
            elif cc == "ES":
                continue
            elif fc.startswith("P.") and (p >= FOREIGN_CITY_MIN_POP or fc == "P.PPLC"):
                n_city += 1
                add_place(esname, "u", "city")
            elif fc == "A.ADM1" and p >= FOREIGN_ADM1_MIN_POP:
                n_adm += 1
                add_place(esname, "u", "state")
    stats.append(f"GeoNames: {n_country} countries and continents, {n_city} foreign cities "
                 f"(>= {FOREIGN_CITY_MIN_POP} inhabitants or capital), {n_adm} first-level "
                 f"regions (>= {FOREIGN_ADM1_MIN_POP})")
    stats.append(f"places: {len(places)} names")

    # ---- write ----
    def write(fname, rows):
        with open(f"{out}/{fname}", "w", encoding="utf-8", newline="\n") as o:
            for cols in sorted(rows, key=lambda r: sort_key(r[0])):
                o.write("\t".join(c for c in cols if c) + "\n")

    def w(n):
        return "word" if n in words else ""

    write("given_names.tsv", [(n, g, w(n)) for n, g in given.items()])
    write("surnames.tsv", [(n, "u", w(n)) for n in surnames])
    write("places.tsv", [(n, g, k, w(n)) for n, (g, k, _) in places.items()])
    for lst, d in (("given", given), ("surname", surnames), ("place", places)):
        stats.append(f"{lst}: {sum(1 for n in d if n in words)} marked word")
    text = "\n".join(stats)
    print(text)
    if stats_path:
        open(stats_path, "w", encoding="utf-8").write(text + "\n")


if __name__ == "__main__":
    a = sys.argv[1:]
    st = None
    if "--stats" in a:
        i = a.index("--stats")
        st = a[i + 1]
        del a[i:i + 2]
    main(*a, stats_path=st)
